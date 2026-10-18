"""Import fixtures are generated locally; no Excel runtime or network is used."""
import csv
import io
import math
import re
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from xml.etree import ElementTree as ET

def cached_values(data, cells):
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    root=ET.fromstring(data)
    for ref,(value,kind) in cells.items():
        cell=root.find(f'.//s:c[@r="{ref}"]',ns)
        if kind: cell.set('t',kind)
        node=cell.find('s:v',ns)
        if node is None:node=ET.SubElement(cell,'{'+ns['s']+'}v')
        node.text=value
    return ET.tostring(root,encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from gui import table_import as importer
from gui.table_import import InputError, inspect_xlsx, read_table


class TableImportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def csv(self, text, encoding='utf-8-sig', name='data.csv'):
        path = self.root / name
        path.write_text(text, encoding=encoding)
        return path

    def workbook(self, rows=None, title='Data'):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title
        for row in rows or []:
            sheet.append(row)
        return workbook, sheet

    def save(self, workbook, name='data.xlsx'):
        path = self.root / name
        workbook.save(path)
        workbook.close()
        return path

    def rewrite(self, path, transform):
        with zipfile.ZipFile(path) as source:
            contents = [(entry.filename, source.read(entry)) for entry in source.infolist()]
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
            for name, data in contents:
                replacement = transform(name, data)
                if replacement is not None:
                    target.writestr(name, replacement)

    def test_csv_types_preserve_text_ids_and_normalize(self):
        path = self.csv('编号;数值;日期;分类\n00123;1.25;2025-01-02;A\n00002;2;2025-02-01;B\n')
        destination = self.root / 'output.csv'
        info = read_table(path, delimiter=';', destination=destination)
        self.assertEqual(info['rows'], 2)
        self.assertEqual(info['types'], {'编号': 'text', '数值': 'numeric', '日期': 'date', '分类': 'text'})
        self.assertEqual(info['sample'][0]['编号'], '00123')
        self.assertTrue(destination.read_bytes().startswith(b'\xef\xbb\xbf'))
        with destination.open(encoding='utf-8-sig', newline='') as handle:
            self.assertEqual(list(csv.reader(handle))[1], ['00123', '1.25', '2025-01-02', 'A'])
        self.assertEqual(info['csv_path'], str(destination))

    def test_gb18030_and_header_selection(self):
        path = self.csv('标题\n姓名,数值\n张三,2\n', encoding='gb18030')
        info = read_table(path, encoding='gb18030', header_row=2)
        self.assertEqual(info['headers'], ['姓名', '数值'])
        self.assertEqual(info['sample'], [{'姓名': '张三', '数值': '2'}])

    def test_ragged_csv_is_rejected_short_and_long(self):
        for body in ['a,b\n1\n', 'a,b\n1,2,3\n', 'a,b\n\n']:
            with self.subTest(body=body), self.assertRaisesRegex(InputError, '表头有 2 列'):
                read_table(self.csv(body))

    def test_invalid_headers(self):
        for header, message in [('a,a', '重复'), ('a, a ', '重复'), ('a,', '空列名'), (' ,b', '空列名'), ('a,\x00', '无效')]:
            with self.subTest(header=header), self.assertRaisesRegex(InputError, message):
                read_table(self.csv(header + '\n1,2\n'))

    def test_invalid_csv_encoding_and_quote(self):
        with self.assertRaisesRegex(InputError, '编码不匹配'):
            read_table(self.csv('名字,数值\n中文,1\n', encoding='gb18030'))
        with self.assertRaisesRegex(InputError, '无法解析'):
            read_table(self.csv('a,b\n"unterminated,2\n'))

    def test_nonfinite_csv_rejected_and_destination_not_replaced(self):
        destination = self.root / 'output.csv'
        destination.write_bytes(b'keep existing file')
        for value in ['NaN', '-nan', 'Infinity', '-inf', '1e999']:
            with self.subTest(value=value), self.assertRaises(InputError):
                read_table(self.csv(f'a,b\n1,2\n3,{value}\n'), destination=destination)
            self.assertEqual(destination.read_bytes(), b'keep existing file')
            self.assertEqual(list(self.root.glob('.table-*')), [])

    def test_source_is_never_overwritten(self):
        path = self.csv('a,b\n1,2\n')
        original = path.read_bytes()
        with self.assertRaisesRegex(InputError, '不能覆盖'):
            read_table(path, destination=path)
        self.assertEqual(path.read_bytes(), original)

    def test_mixed_and_empty_column_types(self):
        info = read_table(self.csv('mixed,empty\n1,\ntext,\n'))
        self.assertEqual(info['types'], {'mixed': 'text', 'empty': 'empty'})
        self.assertEqual(info['columns'][0]['numeric'], 1)
        self.assertEqual(info['columns'][0]['text'], 1)
        self.assertEqual(info['columns'][1]['empty'], 2)

    def test_no_rows_and_bad_arguments(self):
        path = self.csv('a,b\n')
        with self.assertRaisesRegex(InputError, '没有数据行'):
            read_table(path)
        for value in [0, 101, 1.5, '1', True]:
            with self.subTest(value=value), self.assertRaisesRegex(InputError, '1–100'):
                read_table(path, header_row=value)
        with self.assertRaisesRegex(InputError, '不包含工作表'):
            read_table(path, sheet='Data')
        with self.assertRaises(InputError):
            read_table(path, delimiter='|')

    def test_csv_bounds(self):
        with patch.object(importer, 'MAX_ROWS', 1), self.assertRaisesRegex(InputError, '200,000'):
            read_table(self.csv('a\n1\n2\n'))
        with self.assertRaisesRegex(InputError, '200 列'):
            read_table(self.csv(','.join(f'h{x}' for x in range(201)) + '\n' + ','.join('1' for _ in range(201))))
        with patch.object(importer, 'MAX_BYTES', 2), self.assertRaisesRegex(InputError, '上限'):
            read_table(self.csv('a\n1\n'))

    def test_inspection_does_not_validate_readme_header(self):
        workbook, sheet = self.workbook([['Instructions']],'README')
        sheet.merge_cells('A1:D1')
        data = workbook.create_sheet('观测 数据')
        data.append(['Project title'])
        data.append(['id', 'value'])
        data.append(['00123', 3.5])
        path = self.save(workbook)
        info = inspect_xlsx(path)
        self.assertEqual([s['name'] for s in info['sheets']], ['README', '观测 数据'])
        self.assertEqual(info['sheets'][0]['preview'][0][0], 'Instructions')
        self.assertEqual(info['sheets'][0]['merged_ranges'], ['A1:D1'])
        self.assertEqual(info['sheet'], 'README')
        self.assertEqual(info['raw_preview'], info['sheets'][0]['preview'])
        with self.assertRaisesRegex(InputError, '合并单元格'):
            read_table(path)
        selected = read_table(path, sheet='观测 数据', header_row=2)
        self.assertEqual(selected['sample'], [{'id': '00123', 'value': '3.5'}])
        with self.assertRaisesRegex(InputError, '不存在'):
            read_table(path, sheet='观测 数据 ')

    def test_excel_date_and_zero_format_preserved(self):
        workbook, sheet = self.workbook([['id_text', 'id_num', 'date', 'datetime', 'blank', 'boolean'],
                                         ['00123', 123, date(2025, 4, 3), datetime(2025, 4, 3, 15, 30), None, True]])
        sheet['B2'].number_format = '00000'
        sheet['C2'].number_format = 'yyyy-mm-dd'
        sheet['D2'].number_format = 'yyyy-mm-dd hh:mm:ss'
        path = self.save(workbook)
        info = read_table(path, destination=self.root / 'normalized.csv')
        self.assertEqual(info['sample'][0], {'id_text': '00123', 'id_num': '00123', 'date': '2025-04-03',
                                           'datetime': '2025-04-03T15:30:00', 'blank': '', 'boolean': 'TRUE'})
        self.assertEqual(info['types']['id_num'], 'text')
        self.assertEqual(info['types']['date'], 'date')

    def test_formula_caches_errors_and_empty_rows(self):
        workbook, sheet = self.workbook([['id', 'cached', 'uncached', 'error'],
                                         ['A', '=1+2', '=10+2', '#DIV/0!'],
                                         [None, None, None, None], ['B', 5, None, None]])
        path = self.save(workbook)
        self.rewrite(path, lambda name, data: cached_values(data, {'B2':('3',None)})
                     if name == 'xl/worksheets/sheet1.xml' else data)
        info = read_table(path)
        self.assertEqual(info['rows'], 2)
        self.assertEqual(info['sample'][0], {'id': 'A', 'cached': '3', 'uncached': '', 'error': ''})
        self.assertIn('没有可用缓存', ' '.join(info['warnings']))
        self.assertIn('错误单元格', ' '.join(info['warnings']))
        self.assertIn('全空数据行', ' '.join(info['warnings']))
        self.assertIn('缓存可能不是最新', ' '.join(info['warnings']))

    def test_formula_text_result_and_cached_error(self):
        workbook, sheet = self.workbook([['id', 'name', 'error'], ['x', '="00123"', '=1/0']])
        path = self.save(workbook)
        def alter(name, data):
            if name.endswith('sheet1.xml'):
                data = cached_values(data, {'B2':('00123','str'),'C2':('#DIV/0!','e')})
            return data
        self.rewrite(path, alter)
        info = read_table(path)
        self.assertEqual(info['sample'][0]['name'], '00123')
        self.assertEqual(info['sample'][0]['error'], '')

    def test_formula_text_is_never_evaluated(self):
        workbook, sheet = self.workbook([['id', 'formula'], ['x', '=WEBSERVICE("https://example.invalid/data")']])
        path = self.save(workbook)
        info = read_table(path)
        self.assertEqual(info['sample'][0]['formula'], '')
        self.assertIn('不会计算', ' '.join(info['warnings']))

    def test_hidden_sheets_listed(self):
        workbook, sheet = self.workbook([['a'], [1]])
        hidden = workbook.create_sheet('Hidden')
        hidden.append(['b']); hidden.append([2]); hidden.sheet_state = 'hidden'
        path = self.save(workbook)
        self.assertEqual(inspect_xlsx(path)['sheets'][1]['state'], 'hidden')
        self.assertEqual(read_table(path, sheet='Hidden')['rows'], 1)

    def test_lying_dimension_does_not_hide_data_or_expand_rows(self):
        for dimension in [b'A1:A1', b'A1:XFD1048576']:
            workbook, sheet = self.workbook([['a', 'b'], [1, 2], [3, 4]])
            path = self.save(workbook)
            self.rewrite(path, lambda name, data: re.sub(br'<dimension ref="[^"]*"', b'<dimension ref="' + dimension + b'"', data)
                         if name.endswith('sheet1.xml') else data)
            info = read_table(path)
            self.assertEqual(info['rows'], 2)
            self.assertEqual(info['headers'], ['a', 'b'])

    def test_sheet_bounds_still_allow_inspecting_other_sheets(self):
        workbook, sheet = self.workbook([['a'], [1]])
        huge = workbook.create_sheet('Too wide'); huge.cell(1, 201, 'last')
        path = self.save(workbook)
        self.assertFalse(inspect_xlsx(path)['sheets'][1]['selectable'])
        self.assertEqual(read_table(path, sheet='Data')['rows'], 1)
        with self.assertRaisesRegex(InputError, '200 列'):
            read_table(path, sheet='Too wide')
        with patch.object(importer, 'MAX_ROWS', 0), self.assertRaisesRegex(InputError, '200,000'):
            read_table(path, sheet='Data')

    def test_xlsx_duplicate_and_merged_header(self):
        workbook, sheet = self.workbook([['id', 'id'], [1, 2]])
        with self.assertRaisesRegex(InputError, '重复'):
            read_table(self.save(workbook))
        workbook, sheet = self.workbook([['id', None], [1, 2]])
        sheet.merge_cells('A1:B1')
        with self.assertRaisesRegex(InputError, '合并单元格'):
            read_table(self.save(workbook))

    def test_xlsx_nan_infinity_invalid_numeric_xml(self):
        for value in [b'NaN', b'INF', b'-Infinity', b'1e999', b'not-a-number']:
            workbook, sheet = self.workbook([['id', 'n'], ['x', 3]])
            path = self.save(workbook)
            self.rewrite(path, lambda name, data: data.replace(b'<v>3</v>', b'<v>' + value + b'</v>') if name.endswith('sheet1.xml') else data)
            with self.subTest(value=value), self.assertRaisesRegex(InputError, '数值'):
                read_table(path)

    def test_dtd_and_entities_rejected_utf8_and_utf16(self):
        for encoding in ['utf-8', 'utf-16', 'utf-16-be', 'utf-16-le']:
            workbook, sheet = self.workbook([['a'], [1]])
            path = self.save(workbook)
            def alter(name, data):
                if name.endswith('sheet1.xml'):
                    xml = data.decode('utf-8')
                    xml = '<?xml version="1.0" encoding="' + encoding + '"?><!DOCTYPE worksheet [<!ENTITY danger "expanded">]>' + xml
                    return xml.encode(encoding)
                return data
            self.rewrite(path, alter)
            with self.subTest(encoding=encoding), self.assertRaises(InputError):
                inspect_xlsx(path)

    def test_zip_bomb_and_member_count_limits(self):
        workbook, sheet = self.workbook([['a'], [1]])
        path = self.save(workbook)
        with patch.object(importer, 'MAX_UNCOMPRESSED_BYTES', 100), self.assertRaisesRegex(InputError, '150 MB'):
            inspect_xlsx(path)
        with patch.object(importer, 'MAX_ZIP_ENTRIES', 2), self.assertRaisesRegex(InputError, '文件过多'):
            inspect_xlsx(path)

    def test_unsafe_duplicate_and_traversal_members(self):
        for member in ['../escape', '/escape', 'xl\\escape', 'xl/workbook.xml']:
            workbook, sheet = self.workbook([['a'], [1]])
            path = self.save(workbook)
            with zipfile.ZipFile(path, 'a') as archive:
                entry=zipfile.ZipInfo(member)
                entry.filename=member
                archive.writestr(entry, b'invalid')
            with self.subTest(member=member), self.assertRaisesRegex(InputError, '路径重复或无效'):
                inspect_xlsx(path)

    def test_macros_renamed_to_xlsx_rejected(self):
        workbook, sheet = self.workbook([['a'], [1]])
        path = self.save(workbook)
        with zipfile.ZipFile(path, 'a') as archive:
            archive.writestr('xl/vbaProject.bin', b'not executed')
        with self.assertRaisesRegex(InputError, '宏'):
            inspect_xlsx(path)
        workbook, sheet = self.workbook([['a'], [1]])
        path = self.save(workbook)
        self.rewrite(path, lambda name, data: data.replace(importer._WORKBOOK_TYPE.encode(), b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
                     if name == '[Content_Types].xml' else data)
        with self.assertRaisesRegex(InputError, '宏'):
            inspect_xlsx(path)

    def test_external_link_is_not_fetched(self):
        workbook, sheet = self.workbook([['a'], [1]])
        path = self.save(workbook)
        self.rewrite(path, lambda name, data: data.replace(b'</Relationships>',
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/never-fetch" TargetMode="External" Id="rIdExternal"/></Relationships>')
            if name == 'xl/_rels/workbook.xml.rels' else data)
        info = read_table(path)
        self.assertIn('未访问或更新外部内容', ' '.join(info['warnings']))

    def test_inspect_csv_title_rows_and_raw_preview(self):
        path = self.csv('Project title\nid,value\n00123,2\n')
        info = importer.inspect_csv(path)
        self.assertEqual(info['raw_preview'], [['Project title'], ['id', 'value'], ['00123', '2']])
        self.assertEqual(info['rows'], 3)
        self.assertEqual(read_table(path, header_row=2)['rows'], 1)
        with self.assertRaisesRegex(InputError, '无法解析'):
            importer.inspect_csv(self.csv('a,b\n"unclosed,2\n'))

    def test_unique_counts_scan_full_table_and_report_cap(self):
        info = read_table(self.csv('category\nA\nA\nA\nA\nA\nB\nC\n'))
        self.assertEqual(info['columns'][0]['unique_count'], 3)
        self.assertFalse(info['columns'][0]['distinct_truncated'])
        with patch.object(importer, 'UNIQUE_LIMIT', 2):
            info = read_table(self.csv('category\nA\nB\nB\n'))
            self.assertEqual(info['columns'][0]['unique_count'], 2)
            self.assertFalse(info['columns'][0]['distinct_truncated'])
            info = read_table(self.csv('category\nA\nB\nC\nD\n'))
            self.assertEqual(info['columns'][0]['unique_count'], 2)
            self.assertTrue(info['columns'][0]['distinct_truncated'])
            self.assertEqual(info['columns'][0]['unique_limit'], 2)

    def test_workbook_content_type_default_is_supported(self):
        workbook, sheet = self.workbook([['a'], [1]])
        path = self.save(workbook)
        def alter(name, data):
            if name == '[Content_Types].xml':
                data = re.sub(br'<Override PartName="/xl/workbook.xml" ContentType="[^"]*"/>', b'', data)
                data = data.replace(b'Extension="xml" ContentType="application/xml"',
                                    b'Extension="xml" ContentType="' + importer._WORKBOOK_TYPE.encode() + b'"')
            return data
        self.rewrite(path, alter)
        self.assertEqual(read_table(path)['rows'], 1)

    def test_negative_zero_format_and_header_below_merged_title(self):
        workbook, sheet = self.workbook([['Title'], ['id', 'value'], [-3, 1]])
        sheet.merge_cells('A1:D1')
        sheet['A3'].number_format = '00000'
        info = read_table(self.save(workbook), header_row=2)
        self.assertEqual(info['sample'][0]['id'], '-00003')

    def test_demo_only_nonfinite_text_passes_unchanged_with_warning(self):
        path = self.csv('id,value\n1,NaN\n2,Inf\n3,-Infinity\n')
        destination = self.root / 'normalized.csv'
        info = read_table(path, destination=destination, allow_nonfinite_text=True)
        self.assertEqual([row['value'] for row in info['sample']], ['NaN', 'Inf', '-Infinity'])
        self.assertIn('原始计算引擎审计', ' '.join(info['warnings']))
        self.assertIn('NaN', destination.read_text(encoding='utf-8-sig'))
        with self.assertRaises(InputError):
            read_table(path)
        with self.assertRaises(InputError):
            read_table(self.csv('a\n1e999\n'), allow_nonfinite_text=True)
        workbook, sheet = self.workbook([['a'], ['NaN']])
        with self.assertRaises(InputError):
            read_table(self.save(workbook), allow_nonfinite_text=True)


    def test_artifact_tool_templates(self):
        directory = ROOT / 'examples' / 'templates'
        if not (directory / 'gis.xlsx').exists():
            self.skipTest('Optional template deliverables are not built yet')
        for path in [directory / 'gis.xlsx', directory / 'stats.xlsx']:
            with self.subTest(path=path.name):
                inspection = inspect_xlsx(path)
                self.assertGreaterEqual(len(inspection['sheets']), 1)
                info = read_table(path)
                self.assertGreater(info['rows'], 0)


    def test_other_archives_encrypted_like_and_extensions_rejected(self):
        invalid = self.root / 'not-xlsx.xlsx'
        with zipfile.ZipFile(invalid, 'w') as archive:
            archive.writestr('hello.txt', b'hello')
        with self.assertRaisesRegex(InputError, '不是有效'):
            inspect_xlsx(invalid)
        invalid.write_bytes(b'\xd0\xcf\x11\xe0encrypted OLE')
        with self.assertRaisesRegex(InputError, '加密'):
            inspect_xlsx(invalid)
        for name in ['data.xls', 'data.xlsm', 'data.zip']:
            with self.subTest(name=name), self.assertRaisesRegex(InputError, '仅支持'):
                read_table(self.csv('a\n1\n', name=name))


if __name__ == '__main__':
    unittest.main()
