"""Bounded, local-only CSV/XLSX ingestion for the GUI.

``inspect_xlsx`` lists worksheets and their raw first rows, without choosing a
header. ``read_table`` validates a selected table and optionally writes an atomic
UTF-8-BOM/comma CSV snapshot for the unchanged engine. Formula results are only
read from caches; this module never calculates formulas or follows links.
"""
from __future__ import annotations

import csv
import math
import os
import posixpath
import re
import tempfile
import zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from pathlib import Path
from xml.parsers import expat

MAX_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
MAX_ZIP_ENTRIES = 2048
MAX_SHEETS = 200
MAX_ROWS = 200000
MAX_COLUMNS = 200
MAX_HEADER_ROW = 100
PREVIEW_ROWS = 8
SAMPLE_ROWS = 5
UNIQUE_LIMIT = 1001
ENCODINGS = {'utf-8-sig', 'gb18030'}
DELIMITERS = {',', '\t', ';'}
_WORKBOOK_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'
_CELL = re.compile(r'^([A-Za-z]{1,3})([1-9][0-9]*)$')
_NUMBER = re.compile(r'^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?$')
_LEADING_ZERO = re.compile(r'^[+-]?0\d+$')
_NONFINITE = {'nan', '+nan', '-nan', 'inf', '+inf', '-inf', 'infinity', '+infinity', '-infinity'}


class InputError(ValueError):
    """A safe, user-readable table validation error."""


class _Warnings:
    def __init__(self):
        self.counts = Counter()
        self.examples = {}

    def add(self, kind, location=''):
        self.counts[kind] += 1
        examples = self.examples.setdefault(kind, [])
        if location and location not in examples and len(examples) < 3:
            examples.append(location)

    def messages(self):
        labels = {
            'formula': '公式仅读取文件中保存的缓存值，不会计算；缓存可能不是最新结果',
            'formula_empty': '公式没有可用缓存值，已作为空值；请用 Excel / LibreOffice 重新计算并保存后再导入',
            'error': 'Excel 错误单元格已作为空值',
            'external': '工作簿含外部链接，导入时未访问或更新外部内容',
            'blank_row': '全空数据行已跳过',
            'nonfinite_demo': '合成示例中的 NaN / Infinity 文本已原样保留，由原始计算引擎审计',
        }
        result = []
        for kind, count in self.counts.items():
            examples = self.examples.get(kind, [])
            suffix = '（例如 ' + '、'.join(examples) + '）' if examples else ''
            result.append(f'{labels[kind]}：{count} 处{suffix}')
        return result


def _source(path, suffixes):
    path = Path(path)
    if path.suffix.lower() not in suffixes:
        raise InputError('仅支持 .csv 和 .xlsx；不支持 .xls、.xlsm、加密工作簿或其他压缩包')
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise InputError('无法读取数据文件') from exc
    if not path.is_file() or size <= 0 or size > MAX_BYTES:
        raise InputError(f'文件为空或超过当前 {MAX_BYTES//(1024*1024)} MB 上限')
    return path


def _column_number(letters):
    result = 0
    for letter in letters.upper():
        result = result * 26 + ord(letter) - ord('A') + 1
    return result


class _XmlScan:
    """Parse XML before openpyxl, rejecting declarations in any XML encoding."""
    def __init__(self, filename):
        self.filename = filename
        self.content_types = {}
        self.defaults = {}
        self.sheets = []
        self.relationships = {}
        self.external = 0
        self.rows = 0
        self.columns = 0
        self.merges = []
        self.merge_count = 0
        self.merged_header_rows = set()
        self.row = 0
        self.column = 0
        self.cell_type = None
        self.cell_ref = ''
        self.value = None
        self.parser = expat.ParserCreate(namespace_separator='}')
        self.parser.StartDoctypeDeclHandler = self.reject_declaration
        self.parser.EntityDeclHandler = self.reject_declaration
        self.parser.ExternalEntityRefHandler = self.reject_declaration
        self.parser.SetParamEntityParsing(expat.XML_PARAM_ENTITY_PARSING_NEVER)
        self.parser.StartElementHandler = self.start
        self.parser.EndElementHandler = self.end
        self.parser.CharacterDataHandler = self.characters

    def reject_declaration(self, *args):
        raise InputError('XLSX 包含不安全的 XML DTD / 实体声明，已拒绝导入')

    def start(self, name, attrs):
        tag = name.rsplit('}', 1)[-1]
        if self.filename == '[Content_Types].xml' and tag in {'Override', 'Default'}:
            content_type = attrs.get('ContentType', '')
            lowered = content_type.lower()
            if any(x in lowered for x in ('macroenabled', 'vbaproject', 'macrosheet')):
                raise InputError('仅支持不含宏的 .xlsx 工作簿；请另存为 .xlsx')
            if tag == 'Override':
                self.content_types[attrs.get('PartName', '').lstrip('/')] = content_type
            else:
                self.defaults[attrs.get('Extension', '').lower()] = content_type
        if self.filename == 'xl/workbook.xml' and tag == 'sheet':
            rid = next((v for k, v in attrs.items() if k.endswith('}id')), '')
            self.sheets.append({'name': attrs.get('name', ''), 'id': rid, 'state': attrs.get('state', 'visible')})
            if len(self.sheets) > MAX_SHEETS:
                raise InputError(f'工作簿最多支持 {MAX_SHEETS} 个工作表')
        if tag == 'Relationship':
            self.relationships[attrs.get('Id', '')] = attrs
            if attrs.get('TargetMode', '').lower() == 'external':
                self.external += 1
        if tag == 'row':
            try:
                self.row = int(attrs.get('r', self.row + 1))
            except (TypeError, ValueError) as exc:
                raise InputError('XLSX 行号无效') from exc
            self.column = 0
        if tag == 'c':
            self.cell_type = attrs.get('t', 'n')
            self.cell_ref = attrs.get('r', '')
            if self.cell_ref:
                match = _CELL.fullmatch(self.cell_ref)
                if not match:
                    raise InputError('XLSX 单元格位置无效：' + self.cell_ref[:30])
                self.column = _column_number(match.group(1))
                current_row = int(match.group(2))
            else:
                self.column += 1
                current_row = self.row
            self.rows = max(self.rows, current_row)
            self.columns = max(self.columns, self.column)
        elif tag == 'v' and self.cell_type == 'n':
            self.value = ''
        elif tag == 'mergeCell':
            self.merge_count += 1
            bounds = attrs.get('ref', '').split(':')
            matches = [_CELL.fullmatch(bound) for bound in bounds]
            if len(matches) != 2 or not all(matches):
                raise InputError('XLSX 合并单元格范围无效')
            first, last = (int(match.group(2)) for match in matches)
            self.merged_header_rows.update(range(max(1, first), min(last, MAX_HEADER_ROW) + 1))
            if len(self.merges) < 10:
                self.merges.append(attrs.get('ref', ''))

    def characters(self, value):
        if self.value is not None:
            self.value += value
            if len(self.value) > 200:
                raise InputError('XLSX 数值单元格内容过长')

    def end(self, name):
        tag = name.rsplit('}', 1)[-1]
        if tag == 'v' and self.value is not None:
            if self.value.strip():
                try:
                    finite = math.isfinite(float(self.value))
                except (ValueError, OverflowError) as exc:
                    raise InputError('XLSX 数值单元格无效：' + self.cell_ref) from exc
                if not finite:
                    raise InputError('XLSX 含 NaN / Infinity 或超出范围的数值：' + self.cell_ref)
            self.value = None
        elif tag == 'c':
            self.cell_type = None


def _preflight(path):
    """Validate every ZIP member and XML document without extracting files."""
    scans = {}
    warnings = _Warnings()
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise InputError(f'XLSX 内部文件过多（最多 {MAX_ZIP_ENTRIES} 个）')
            names = set()
            total = 0
            for entry in entries:
                name = entry.filename
                # Windows ZipInfo normalizes backslashes and truncates NULs.
                # Validate the archive's original spelling before that normalization.
                if entry.orig_filename != name or '\\' in entry.orig_filename or '\x00' in entry.orig_filename:
                    raise InputError('XLSX 内部文件路径重复或无效')
                if (name in names or name.startswith('/') or '\\' in name
                        or '..' in name.split('/') or '\x00' in name):
                    raise InputError('XLSX 内部文件路径重复或无效')
                names.add(name)
                if entry.flag_bits & 1:
                    raise InputError('不支持加密工作簿，请先移除密码并另存为 .xlsx')
                if entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                    raise InputError('XLSX 使用不支持的压缩方式')
                if any(part in name.lower() for part in ('vbaproject', 'vbadata', '/macrosheets/')):
                    raise InputError('仅支持不含宏的 .xlsx 工作簿')
                total += entry.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise InputError('XLSX 解压后超过 150 MB 安全限制')
            if not {'[Content_Types].xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels'}.issubset(names):
                raise InputError('文件不是有效的 .xlsx 工作簿')
            ordered = sorted(entries, key=lambda e: e.filename != '[Content_Types].xml')
            actual = 0
            content_types, defaults = {}, {}
            for entry in ordered:
                name = entry.filename
                content_type = content_types.get(name, defaults.get(name.rsplit('.', 1)[-1].lower(), ''))
                is_xml = name.lower().endswith(('.xml', '.rels', '.vml')) or 'xml' in content_type.lower()
                scan = _XmlScan(name) if is_xml else None
                with archive.open(entry) as member:
                    while True:
                        chunk = member.read(64 * 1024)
                        if not chunk:
                            break
                        actual += len(chunk)
                        if actual > MAX_UNCOMPRESSED_BYTES:
                            raise InputError('XLSX 解压后超过 150 MB 安全限制')
                        if scan:
                            scan.parser.Parse(chunk, False)
                if scan:
                    scan.parser.Parse(b'', True)
                    scans[name] = scan
                    for _ in range(scan.external):
                        warnings.add('external')
                    if name == '[Content_Types].xml':
                        content_types, defaults = scan.content_types, scan.defaults
            if content_types.get('xl/workbook.xml', defaults.get('xml')) != _WORKBOOK_TYPE:
                raise InputError('仅支持标准 .xlsx 工作簿；不支持含宏、模板或其他类型的文件')
    except InputError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError, EOFError, ValueError, expat.ExpatError, NotImplementedError) as exc:
        raise InputError('XLSX 已损坏、加密或不是有效工作簿；请重新另存为 .xlsx') from exc
    rels = scans['xl/_rels/workbook.xml.rels'].relationships
    sheets = []
    for sheet in scans['xl/workbook.xml'].sheets:
        rel = rels.get(sheet['id'], {})
        if not rel.get('Type', '').endswith('/worksheet'):
            continue  # Chartsheets do not contain importable tables.
        if rel.get('TargetMode', '').lower() == 'external':
            raise InputError('不支持指向外部文件的工作表')
        target = rel.get('Target', '')
        target = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('xl', target))
        scan = scans.get(target)
        if scan is None:
            raise InputError('XLSX 工作表缺失或格式无效')
        sheets.append({**sheet, 'rows': scan.rows, 'columns': scan.columns,
                       'merged_ranges': scan.merges, 'merged_count': scan.merge_count,
                       'merged_header_rows': sorted(scan.merged_header_rows)})
    if not sheets:
        raise InputError('XLSX 没有可读取的工作表')
    return sheets, warnings


@contextmanager
def _workbooks(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise InputError('Excel 导入需要 openpyxl，请先运行软件附带的依赖安装程序') from exc
    values = formulas = None
    try:
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        yield values, formulas
    except InputError:
        raise
    except Exception as exc:
        raise InputError('XLSX 无法读取；请检查工作簿内容并重新另存为 .xlsx') from exc
    finally:
        if values is not None:
            values.close()
        if formulas is not None:
            formulas.close()


def _string(value, location='', allow_nonfinite_text=False, warnings=None):
    if value is None:
        return '', 'empty'
    if isinstance(value, bool):
        return ('TRUE' if value else 'FALSE'), 'text'
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(), 'date'
    if isinstance(value, timedelta):
        return str(value.total_seconds()), 'numeric'
    if isinstance(value, (int, float)):
        try:
            finite = math.isfinite(value)
        except OverflowError:
            finite = False
        if not finite:
            raise InputError('数据含 NaN / Infinity 或超出范围的数值：' + location)
        return str(value), 'numeric'
    text = str(value)
    if '\x00' in text:
        raise InputError('数据包含无效的空字符：' + location)
    clean = text.strip()
    if not clean:
        return text, 'empty'
    if clean.lower() in _NONFINITE:
        if allow_nonfinite_text:
            if warnings is not None:
                warnings.add('nonfinite_demo', location)
            return text, 'text'
        raise InputError('数据含 NaN / Infinity：' + location + '；请改为空值或有效数值')
    if _NUMBER.fullmatch(clean):
        try:
            finite = math.isfinite(float(clean))
        except (ValueError, OverflowError):
            finite = False
        if not finite:
            raise InputError('数据含超出范围的数值：' + location)
        if not _LEADING_ZERO.fullmatch(clean):
            return text, 'numeric'
    if re.match(r'^\d{4}-\d{2}-\d{2}(?:$|[T ])', clean):
        try:
            datetime.fromisoformat(clean)
        except ValueError:
            pass
        else:
            return text, 'date'
    return text, 'text'


def _cell(value_cell, formula_cell, warnings, location):
    formula = getattr(formula_cell, 'data_type', None) == 'f'
    value = value_cell.value
    if formula:
        warnings.add('formula', location)
        if value is None:
            warnings.add('formula_empty', location)
    if getattr(value_cell, 'data_type', None) == 'e':
        warnings.add('error', location)
        return '', 'empty'
    # Numeric cells formatted 00000 are commonly identifiers. Preserve the
    # displayed leading zeroes without assigning any logical role to the column.
    fmt = getattr(formula_cell, 'number_format', 'General')
    if (isinstance(value, (int, float)) and not isinstance(value, bool)
            and re.fullmatch(r'0{2,}', fmt or '') and math.isfinite(value) and float(value).is_integer()):
        integer = int(value)
        return ('-' if integer < 0 else '') + str(abs(integer)).zfill(len(fmt)), 'text'
    if isinstance(value, datetime):
        from openpyxl.styles.numbers import is_datetime
        if is_datetime(fmt) == 'date':
            return value.date().isoformat(), 'date'
    return _string(value, location)


def _worksheet_rows(values, formulas, sheet, start, end, columns, warnings):
    value_rows = values[sheet].iter_rows(min_row=start, max_row=end, max_col=columns)
    formula_rows = formulas[sheet].iter_rows(min_row=start, max_row=end, max_col=columns)
    for row_number, (vrow, frow) in enumerate(zip(value_rows, formula_rows), start):
        strings, kinds = [], []
        for index, (vcell, fcell) in enumerate(zip(vrow, frow), 1):
            location = sheet + '!' + (getattr(fcell, 'coordinate', '') or f'行 {row_number} 列 {index}')
            text, kind = _cell(vcell, fcell, warnings, location)
            strings.append(text)
            kinds.append(kind)
        yield row_number, strings, kinds


def inspect_xlsx(path):
    """Return sheet metadata and raw first rows; no header is assumed or validated."""
    path = _source(path, {'.xlsx'})
    metadata, warnings = _preflight(path)
    result = []
    with _workbooks(path) as (values, formulas):
        for item in metadata:
            sheet = {k: v for k, v in item.items() if k != 'id'}
            sheet['selectable'] = item['columns'] <= MAX_COLUMNS and item['rows'] <= MAX_ROWS + MAX_HEADER_ROW
            sheet['preview'] = []
            if item['rows'] and item['columns']:
                for _, strings, _ in _worksheet_rows(values, formulas, item['name'], 1,
                        min(PREVIEW_ROWS, item['rows']), min(MAX_COLUMNS, item['columns']), warnings):
                    sheet['preview'].append([value[:200] for value in strings])
            result.append(sheet)
    return {'format': 'xlsx', 'sheets': result, 'warnings': warnings.messages(),
            'raw_preview': result[0]['preview'], 'sheet': result[0]['name'], 'header_row': 1,
            'preview_rows': PREVIEW_ROWS, 'max_header_row': MAX_HEADER_ROW}


def inspect_csv(path, encoding='utf-8-sig', delimiter=','):
    """Inspect CSV records without assuming that the first record is a header.

    All records are parsed and bounded. Ragged records remain visible for header
    selection, and read_table always performs strict selected-table validation.
    """
    path = _source(path, {'.csv'})
    if encoding not in ENCODINGS or delimiter not in DELIMITERS:
        raise InputError('请选择 UTF-8 或 GB18030 编码，以及逗号、制表符或分号分隔符')
    preview, count, columns = [], 0, 0
    try:
        with path.open('r', encoding=encoding, newline='') as handle:
            reader = csv.reader(handle, delimiter=delimiter, strict=True)
            for values in reader:
                count += 1
                if count > MAX_ROWS + MAX_HEADER_ROW:
                    raise InputError('CSV 超过 200,000 个数据行及 100 个表头前置行；请分批导入')
                if len(values) > MAX_COLUMNS:
                    raise InputError('CSV 超过 200 列，请检查分隔符或整理字段')
                columns = max(columns, len(values))
                strings = [_string(value, f'第 {reader.line_num} 行')[0] for value in values]
                if len(preview) < PREVIEW_ROWS:
                    preview.append([value[:200] for value in strings])
    except UnicodeError as exc:
        raise InputError('CSV 编码不匹配，请尝试 GB18030；或另存为 UTF-8 CSV') from exc
    except csv.Error as exc:
        raise InputError('CSV 无法解析：' + str(exc)) from exc
    return {'format': 'csv', 'sheets': [], 'sheet': None, 'header_row': 1, 'raw_preview': preview,
            'rows': count, 'columns': columns, 'warnings': [], 'preview_rows': PREVIEW_ROWS,
            'max_header_row': MAX_HEADER_ROW}


def _headers(values):
    if not values or not any(value.strip() for value in values):
        raise InputError('所选表头行为空，请选择正确的表头行')
    if len(values) > MAX_COLUMNS:
        raise InputError('界面最多支持 200 列，请整理所需字段后重试')
    if any(not value.strip() or '\x00' in value for value in values):
        raise InputError('表头包含空列名或无效字符；合并表头请先取消合并并补全列名，或选择正确的表头行')
    normalized = [value.strip() for value in values]
    repeated = [name for name, count in Counter(normalized).items() if count > 1]
    if repeated:
        raise InputError('表头存在重复列名：' + '、'.join(repeated[:5]) + '；请修改后重试')
    return values


class _Summary:
    def __init__(self, headers):
        self.headers = headers
        self.rows = 0
        self.sample = []
        self.counts = [Counter() for _ in headers]
        self.examples = [[] for _ in headers]
        self.unique = [set() for _ in headers]
        self.unique_truncated = [False for _ in headers]

    def add(self, values, kinds):
        self.rows += 1
        if self.rows > MAX_ROWS:
            raise InputError('界面最多支持 200,000 行；未抽样，请分批或使用原命令行入口')
        if len(self.sample) < SAMPLE_ROWS:
            self.sample.append({key: value[:200] for key, value in zip(self.headers, values)})
        for index, (value, kind) in enumerate(zip(values, kinds)):
            self.counts[index][kind] += 1
            if kind != 'empty' and value not in self.unique[index]:
                if len(self.unique[index]) < UNIQUE_LIMIT:
                    self.unique[index].add(value)
                else:
                    self.unique_truncated[index] = True
            if kind != 'empty' and value not in self.examples[index] and len(self.examples[index]) < 3:
                self.examples[index].append(value[:200])

    def result(self):
        if not self.rows:
            raise InputError('表格只有表头，没有数据行')
        types, columns = {}, []
        for index, name in enumerate(self.headers):
            counts = self.counts[index]
            kinds = set(counts) - {'empty'}
            suggested = next(iter(kinds)) if len(kinds) == 1 else ('text' if kinds else 'empty')
            types[name] = suggested
            columns.append({'name': name, 'type': suggested, 'suggested_type': suggested,
                            'nonempty': self.rows - counts['empty'], 'empty': counts['empty'],
                            'numeric': counts['numeric'], 'text': counts['text'], 'date': counts['date'],
                            'examples': self.examples[index], 'unique_count': len(self.unique[index]),
                            'unique_limit': UNIQUE_LIMIT, 'distinct_truncated': self.unique_truncated[index]})
        return {'headers': self.headers, 'rows': self.rows, 'sample': self.sample,
                'types': types, 'columns': columns}


@contextmanager
def _output(destination, source):
    if destination is None:
        yield None
        return
    target = Path(destination)
    if target.resolve() == source.resolve():
        raise InputError('标准化 CSV 必须保存到不同的文件，不能覆盖原始数据')
    target.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(mode='w', encoding='utf-8-sig', newline='',
                                         dir=target.parent, prefix='.table-', suffix='.csv', delete=False)
    temporary = Path(handle.name)
    try:
        with handle:
            yield csv.writer(handle, delimiter=',', lineterminator='\n')
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def read_table(path, encoding='utf-8-sig', delimiter=',', sheet=None, header_row=1, destination=None,
               allow_nonfinite_text=False):
    """Validate CSV or one XLSX worksheet, optionally normalizing to a CSV file.

    The return value is JSON-safe: ``headers``, ``rows`` (data-row count),
    ``sample`` (five row dictionaries), ``types`` (column-to-type mapping),
    ``columns`` (type counts/examples), ``warnings``, ``format``, ``sheet``, and
    ``header_row``. ``csv_path`` is added only when destination was written.
    Types describe observed values only, never coordinates or other field roles.
    The explicit demo-only allow_nonfinite_text option preserves literal CSV
    NaN/Inf audit fixtures with a warning; it does not relax XLSX validation.
    """
    path = _source(path, {'.csv', '.xlsx'})
    if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= MAX_HEADER_ROW:
        raise InputError('表头行必须是 1–100 之间的整数')
    warnings = _Warnings()
    result = None
    with _output(destination, path) as writer:
        if path.suffix.lower() == '.csv':
            if encoding not in ENCODINGS or delimiter not in DELIMITERS:
                raise InputError('请选择 UTF-8 或 GB18030 编码，以及逗号、制表符或分号分隔符')
            if sheet is not None:
                raise InputError('CSV 不包含工作表，请清除工作表选择')
            try:
                with path.open('r', encoding=encoding, newline='') as handle:
                    reader = csv.reader(handle, delimiter=delimiter, strict=True)
                    header = None
                    for _ in range(header_row):
                        header = next(reader, None)
                    headers = _headers(header or [])
                    summary = _Summary(headers)
                    if writer:
                        writer.writerow(headers)
                    for values in reader:
                        if len(values) != len(headers):
                            raise InputError(f'CSV 第 {reader.line_num} 行有 {len(values)} 列，表头有 {len(headers)} 列；请修正缺列或多列')
                        converted = [_string(value, f'第 {reader.line_num} 行 / {headers[index]}',
                                             allow_nonfinite_text=allow_nonfinite_text, warnings=warnings)
                                     for index, value in enumerate(values)]
                        strings, kinds = zip(*converted)
                        summary.add(strings, kinds)
                        if writer:
                            writer.writerow(strings)
                    result = summary.result()
            except UnicodeError as exc:
                raise InputError('CSV 编码不匹配，请尝试 GB18030；或另存为 UTF-8 CSV') from exc
            except csv.Error as exc:
                raise InputError('CSV 无法解析：' + str(exc)) from exc
        else:
            metadata, warnings = _preflight(path)
            if sheet is None:
                sheet = metadata[0]['name']
            if not isinstance(sheet, str):
                raise InputError('请选择工作表名称')
            item = next((item for item in metadata if item['name'] == sheet), None)
            if item is None:
                raise InputError('所选工作表不存在：' + sheet[:100])
            if item['columns'] > MAX_COLUMNS:
                raise InputError('所选工作表超过 200 列，请整理所需字段后重试')
            if item['rows'] - header_row > MAX_ROWS:
                raise InputError('所选工作表超过 200,000 个数据行；未抽样，请分批导入')
            if header_row in item['merged_header_rows']:
                raise InputError('表头包含合并单元格，请取消合并并补全列名，或选择正确的表头行')
            if item['rows'] < header_row or not item['columns']:
                raise InputError('所选表头行为空，请选择正确的表头行')
            with _workbooks(path) as (values, formulas):
                rows = _worksheet_rows(values, formulas, sheet, header_row, item['rows'], item['columns'], warnings)
                _, header, _ = next(rows)
                headers = _headers(header)
                summary = _Summary(headers)
                if writer:
                    writer.writerow(headers)
                for row_number, strings, kinds in rows:
                    if all(kind == 'empty' for kind in kinds):
                        warnings.add('blank_row', f'{sheet}!行 {row_number}')
                        continue
                    summary.add(strings, kinds)
                    if writer:
                        writer.writerow(strings)
                result = summary.result()
        result.update(format=path.suffix.lower()[1:], sheet=sheet, header_row=header_row,
                      warnings=warnings.messages())
    if destination is not None:
        result['csv_path'] = str(Path(destination))
    return result
